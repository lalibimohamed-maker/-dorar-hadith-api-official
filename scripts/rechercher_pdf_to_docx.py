#!/usr/bin/env python3
import argparse, hashlib, json, os, re, subprocess, sys, unicodedata, zipfile
from pathlib import Path
import xml.etree.ElementTree as ET
import pymupdf
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from bidi.algorithm import get_display
import arabic_reshaper

ARABIC_RE=re.compile(r"[\u0600-\u06ff\u0750-\u077f\u08a0-\u08ff\ufb50-\ufdff\ufe70-\ufeff]")
PUA_RE=re.compile(r"[\ue000-\uf8ff\U000f0000-\U000ffffd\U00100000-\U0010fffd]")
MOJIBAKE_RE=re.compile(r"(?:Ã.|Ø.|Ù.|â€|ï»¿|�)")
MAX_INPUT_MB=int(os.environ.get("RECHERCHER_PDF_DOCX_MAX_MB","50"))
KNOWN_ISLAMIC_LIGATURES={"ﷲ":"ﷲ","ﷳ":"ﷳ","ﷴ":"ﷴ","ﷵ":"ﷵ","ﷶ":"ﷶ","ﷷ":"ﷷ","ﷸ":"ﷸ","ﷹ":"ﷹ","ﷺ":"ﷺ","ﷻ":"ﷻ","﷼":"﷼","﷽":"﷽"}
INVALID_XML_RE=re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
ARABIC_STOP={"و","في","من","على","إلى","عن","أن","إن","ما","لا","لم","لن","هو","هي","هذا","هذه","ذلك","التي","الذي","ثم"}
ARABIC_HINTS={"الله","رسول","محمد","قال","عن","حدثنا","أخبرنا","كتاب","باب","قالت","صلى","عليه","وسلم","تعالى","الحديث","القرآن"}

def sha256(p):
    h=hashlib.sha256()
    with open(p,"rb") as f:
        for b in iter(lambda:f.read(1048576),b""): h.update(b)
    return h.hexdigest()

def sanitize_xml_text(s):
    return INVALID_XML_RE.sub("",s.replace("\r\n","\n").replace("\r","\n"))

def protect_symbols(s):
    return "".join(KNOWN_ISLAMIC_LIGATURES.get(c,c) for c in unicodedata.normalize("NFC",s))

def pua_info(s):
    vals=sorted(set(ord(c) for c in s if PUA_RE.match(c)))
    return {"detected":bool(vals),"codepoints":[f"U+{v:04X}" for v in vals],"unknown":bool(vals)}

def norm_for_compare(s):
    s=protect_symbols(s)
    s=unicodedata.normalize("NFC",s)
    s=re.sub(r"\s+"," ",s).strip()
    return s

def levenshtein_ratio(a,b):
    a,b=norm_for_compare(a),norm_for_compare(b)
    if a==b: return 0.0
    if not a: return 1.0 if b else 0.0
    if not b: return 1.0
    if len(a)>len(b): a,b=b,a
    prev=list(range(len(a)+1))
    for j,cb in enumerate(b,1):
        cur=[j]
        for i,ca in enumerate(a,1):
            cur.append(min(cur[-1]+1,prev[i]+1,prev[i-1]+(ca!=cb)))
        prev=cur
    return round(prev[-1]/max(len(a),len(b)),6)

def token_metrics(a,b):
    def toks(s):
        return [x for x in re.findall(r"[\w\u0600-\u06ff]+",norm_for_compare(s)) if x not in ARABIC_STOP]
    aa,bb=toks(a),toks(b)
    aset,bset=set(aa),set(bb)
    return {"source_tokens":len(aa),"derived_tokens":len(bb),
            "preserved_unique_token_ratio":round(len(aset&bset)/max(1,len(aset)),6)}

def visual_order_candidate(text):
    if not ARABIC_RE.search(text): return {"applied":False,"confidence":0.0,"reason":"non-arabic"}
    logical=get_display(text)
    if logical==text: return {"applied":False,"confidence":0.0,"reason":"already-logical-or-ambiguous"}
    def score(s):
        words=re.findall(r"[\u0600-\u06ff]{2,}",s)
        return sum(w in ARABIC_HINTS for w in words)*3 + sum(1 for w in words if len(w)>=3)
    shaped=arabic_reshaper.reshape(text)
    shaped_logical=arabic_reshaper.reshape(logical)
    before=score(text); after=score(logical)
    shaping_delta=(shaped_logical.count("ﻻ")+shaped_logical.count("ﷲ"))-(shaped.count("ﻻ")+shaped.count("ﷲ"))
    conf=0.0
    if after>before: conf=0.88
    elif after==before and shaping_delta>0: conf=0.72
    if conf>=0.85:
        return {"applied":True,"confidence":conf,"reason":"high-confidence visual-order recovery","text":logical}
    return {"applied":False,"confidence":conf,"reason":"ambiguous; preserved logical text"}

def set_rtl(p,v):
    pPr=p._p.get_or_add_pPr(); old=pPr.find(qn("w:bidi"))
    if v and old is None:
        el=OxmlElement("w:bidi"); el.set(qn("w:val"),"1"); pPr.append(el)
    elif not v and old is not None: pPr.remove(old)

def set_font(r,name):
    rPr=r._r.get_or_add_rPr(); rf=rPr.find(qn("w:rFonts"))
    if rf is None: rf=OxmlElement("w:rFonts"); rPr.append(rf)
    for a in ("ascii","hAnsi","cs","eastAsia"): rf.set(qn("w:"+a),name)

def add_text(d,text):
    text=sanitize_xml_text(protect_symbols(text))
    if not text.strip(): return
    p=d.add_paragraph(); ar=bool(ARABIC_RE.search(text)); set_rtl(p,ar)
    r=p.add_run(text); set_font(r,"Amiri" if ar else "Liberation Serif")

def bbox_overlap(a,b):
    ax0,ay0,ax1,ay1=a; bx0,by0,bx1,by1=b
    ix=max(0,min(ax1,bx1)-max(ax0,bx0)); iy=max(0,min(ay1,by1)-max(ay0,by0))
    return ix*iy/max(1,(ax1-ax0)*(ay1-ay0))

def ocr_image_block(page,bbox):
    pix=page.get_pixmap(matrix=pymupdf.Matrix(2,2),clip=pymupdf.Rect(*bbox),alpha=False)
    import tempfile
    with tempfile.TemporaryDirectory(prefix="rechercher-ocr-") as td:
        img=Path(td)/"region.png"; pix.save(img)
        cmd=["tesseract",str(img),"stdout","-l",os.environ.get("RECHERCHER_OCR_LANG","ara"),"--psm","6"]
        try:
            p=subprocess.run(cmd,text=True,capture_output=True,check=True,timeout=180)
            return p.stdout.strip()
        except Exception:
            return ""

def page_regions(page):
    data=page.get_text("dict")
    blocks=data.get("blocks",[])
    text_blocks=[]; image_blocks=[]
    for b in blocks:
        bbox=tuple(b.get("bbox",(0,0,0,0)))
        if b.get("type")==0:
            text="\n".join("".join(span.get("text","") for span in line.get("spans",[])) for line in b.get("lines",[])).strip()
            if text: text_blocks.append({"bbox":bbox,"text":text,"kind":"digital"})
        elif b.get("type")==1: image_blocks.append({"bbox":bbox,"kind":"image"})
    out=list(text_blocks); ocr_count=0
    for ib in image_blocks:
        if any(bbox_overlap(ib["bbox"],tb["bbox"])>0.65 for tb in text_blocks): continue
        t=ocr_image_block(page,ib["bbox"])
        if t:
            out.append({"bbox":ib["bbox"],"text":t,"kind":"ocr"}); ocr_count+=1
    out.sort(key=lambda x:(round(x["bbox"][1]/6),-x["bbox"][0]))
    return out,ocr_count,len(text_blocks),len(image_blocks)

def validate_docx(path):
    if not zipfile.is_zipfile(path): return False
    with zipfile.ZipFile(path) as z:
        names=set(z.namelist())
        required={"[Content_Types].xml","word/document.xml"}
        if not required.issubset(names): return False
        for n in required:
            ET.fromstring(z.read(n))
    return True

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--input",required=True); ap.add_argument("--output",required=True); ap.add_argument("--manifest",required=True); ap.add_argument("--strict",action="store_true")
    a=ap.parse_args(); inp=Path(a.input); out=Path(a.output); out.mkdir(parents=True,exist_ok=True); entries=[]
    for pdf in sorted(inp.rglob("*.pdf")):
        e={"source_pdf":str(pdf),"status":"failed","content_policy":"preserve-verbatim","derivation":"pdf-structural-extraction","review_status":"review-required","sacred_text_flag":"requires_review","arabic_alignment_verified":False,"ocr_used":False,"text_repair_applied":False}
        try:
            if pdf.stat().st_size>MAX_INPUT_MB*1048576:
                e.update(status="deferred-large-file",reason=f"input exceeds {MAX_INPUT_MB} MiB"); entries.append(e); continue
            e["source_pdf_sha256"]=sha256(pdf); doc=pymupdf.open(pdf); d=Document(); all_source=[]; all_derived=[]; all_derived_digital=[]; pstats=[]
            for page_no,page in enumerate(doc,1):
                regions,oc,dc,ic=page_regions(page); page_source="\n".join(r["text"] for r in regions if r["kind"]=="digital"); page_derived=[]; page_derived_digital=[]
                for r in regions:
                    original=r["text"]; repair=visual_order_candidate(original); text=repair.get("text",original) if repair["applied"] else original
                    text=sanitize_xml_text(protect_symbols(text)); page_derived.append(text);\n                    if r["kind"]=="digital": page_derived_digital.append(text)\n                    add_text(d,text)
                    if repair["applied"]: e["text_repair_applied"]=True
                    pi=pua_info(original)
                    if pi["detected"]: e.setdefault("pua",{"detected":True,"codepoints":[]}); e["pua"]["codepoints"]=sorted(set(e["pua"]["codepoints"]+pi["codepoints"]))
                all_source.append(page_source); all_derived.append("\n".join(page_derived)); all_derived_digital.append("\n".join(page_derived_digital))
                pstats.append({"page":page_no,"digital_blocks":dc,"image_blocks":ic,"ocr_blocks":oc,"ocr_used":oc>0})
                e["ocr_used"]=e["ocr_used"] or oc>0
            source_digital="\n".join(all_source); derived_text="\n".join(all_derived); derived_digital="\n".join(all_derived_digital)
            digital_chars=len(source_digital); derived_chars=len(derived_text)
            image_pages=sum(1 for x in pstats if x["image_blocks"] and x["digital_blocks"]==0)
            hybrid_pages=sum(1 for x in pstats if x["image_blocks"] and x["digital_blocks"])
            if digital_chars==0 and image_pages: kind="scanned"
            elif hybrid_pages or (image_pages and digital_chars): kind="hybrid"
            else: kind="digital-native"
            e.update(pdf_kind=kind,source_char_count=digital_chars,derived_char_count=derived_chars,derived_digital_char_count=len(derived_digital),page_statistics=pstats)
            if digital_chars:
                e["loss_ratio"]=levenshtein_ratio(source_digital,derived_digital)
                e["comparison_mode"]="normalized-levenshtein"
                e["token_metrics"]=token_metrics(source_digital,derived_digital)
            else:
                e["loss_ratio"]=None; e["comparison_mode"]="not-comparable-ocr-only"
            e["mojibake_detected"]=bool(MOJIBAKE_RE.search(derived_text))
            if e["mojibake_detected"]: e["review_status"]="review-required"
            if e.get("pua",{}).get("detected"): e["review_status"]="review-required"
            target=out/(pdf.relative_to(inp).with_suffix(".docx")); target.parent.mkdir(parents=True,exist_ok=True); d.save(target)
            e["derived_docx"]=str(target); e["derived_docx_sha256"]=sha256(target); e["arabic_alignment_verified"]=bool(digital_chars and not e["text_repair_applied"] and not e.get("mojibake_detected",False))
            e["quality_validation"]={"docx_package":validate_docx(target),"loss_ratio":e["loss_ratio"],"paragraphs":len(d.paragraphs),"xml_safe":validate_docx(target)}
            if not e["quality_validation"]["docx_package"]: raise ValueError("DOCX package/XML validation failed")
            e["status"]="converted"; doc.close()
        except Exception as x: e["error"]=str(x)
        entries.append(e)
    m={"schema":"rechercher/pdf-to-docx/v3","converter":"PyMuPDF+python-docx","policy":"derived-only; source PDFs are never modified or deleted","max_input_mb":MAX_INPUT_MB,"visual_order":"conservative python-bidi detector; arabic-reshaper used only for detection scoring","loss_metric":"normalized Levenshtein distance; OCR-only files are not assigned a fabricated loss ratio","hybrid_policy":"digital text blocks first; OCR only non-overlapping image regions; merge by page coordinates","pua_policy":"known Islamic ligatures preserved; unknown PUA is retained and marked review-required","files":entries,"total_pdfs":len(entries),"converted":sum(e["status"]=="converted" for e in entries),"deferred_large_files":sum(e["status"]=="deferred-large-file" for e in entries),"failed":sum(e["status"]=="failed" for e in entries)}
    Path(a.manifest).write_text(json.dumps(m,ensure_ascii=False,indent=2)+"\n",encoding="utf-8"); print(json.dumps({k:m[k] for k in ("total_pdfs","converted","deferred_large_files","failed")},ensure_ascii=False)); return 1 if a.strict and m["failed"] else 0
if __name__=="__main__": sys.exit(main())
